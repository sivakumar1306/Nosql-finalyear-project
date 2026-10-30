#!/usr/bin/env python3
"""
Build the IEEE-style project paper (two-column Word document).

    python docs/report/build_ieee_report.py
    soffice --headless --convert-to pdf docs/report/IEEE_Report.docx --outdir docs/report

Figures are read from docs/report/figures (see make_figures.py and the
evaluation scripts for how the numbers and charts were produced).
"""

import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")
OUT = os.path.join(HERE, "IEEE_Report.docx")

TITLE = ("Hybrid Rule and Isolation Forest Anomaly Detection for Microservice Logs on MongoDB "
         "with Evidence-Grounded LLM Incident Investigation")
# (name, role line, email or None)
AUTHORS = [
    ("G Sivakumar", "Reg. No. 23BCE1834", "sivakumar.g2023@vitstudent.ac.in"),
    ("Afzal Ahamed A", "Reg. No. 23BCE1353", "afzalahamed.a2023@vitstudent.ac.in"),
    ("Manjula V", "Associate Professor Senior (Guide)", None),
]
AFFILIATION = ["School of Computer Science and Engineering", "Vellore Institute of Technology",
               "Chennai, India"]

COLUMN_WIDTH = Inches(3.45)
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]


MATH_TOKEN = re.compile(r"\b([A-Za-z])_([a-z])\b|\^\(([^)]*)\)")


class Paper:
    def __init__(self):
        self.doc = Document()
        self.section_no = 0
        self.sub_no = 0
        self.figure_no = 0
        self.table_no = 0
        self.equation_no = 0
        self._setup()

    # ---------- layout ----------
    def _setup(self):
        style = self.doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(10)
        fmt = style.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0

        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1.0)
        sec.left_margin = sec.right_margin = Inches(0.625)

    def start_two_columns(self):
        sec = self.doc.add_section(WD_SECTION.CONTINUOUS)
        cols = sec._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sec._sectPr.append(cols)
        cols.set(qn("w:num"), "2")
        cols.set(qn("w:space"), "360")

    @staticmethod
    def _run(paragraph, text, size=None, bold=False, italic=False, small_caps=False):
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if small_caps:
            run.font.small_caps = True
        return run

    # ---------- front matter ----------
    def title_block(self):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        self._run(p, TITLE, size=22)

        table = self.doc.add_table(rows=1, cols=len(AUTHORS))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, (name, role, email) in zip(table.rows[0].cells, AUTHORS):
            first = cell.paragraphs[0]
            first.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._run(first, name, size=11)
            lines = [(role, False, 9.5)] + [(a, True, 9.5) for a in AFFILIATION]
            if email:
                lines.append((email, False, 9))
            for line, italic, size in lines:
                para = cell.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._run(para, line, size=size, italic=italic)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)

    def abstract(self, text, keywords):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.125)
        self._run(p, "Abstract", size=9, bold=True, italic=True)
        self._run(p, "—" + text, size=9, bold=True)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.125)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        self._run(p, "Index Terms", size=9, bold=True, italic=True)
        self._run(p, "—" + keywords, size=9, bold=True)

    # ---------- body ----------
    def _math_runs(self, paragraph, text, size=None, bold=False, italic=False):
        """Write text, turning x_i into x with subscript i and ^(...) into superscript."""
        pos = 0
        for match in MATH_TOKEN.finditer(text):
            if match.start() > pos:
                self._run(paragraph, text[pos:match.start()], size, bold, italic)
            if match.group(1):
                self._run(paragraph, match.group(1), size, bold, True)
                sub = self._run(paragraph, match.group(2), size, bold, True)
                sub.font.subscript = True
            else:
                sup = self._run(paragraph, match.group(3), size, bold, italic)
                sup.font.superscript = True
            pos = match.end()
        if pos < len(text):
            self._run(paragraph, text[pos:], size, bold, italic)

    def heading(self, text):
        self.section_no += 1
        self.sub_no = 0
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        self._run(p, f"{ROMAN[self.section_no - 1]}. {text}", small_caps=True)

    def subheading(self, text):
        self.sub_no += 1
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        self._run(p, f"{chr(64 + self.sub_no)}. {text}", italic=True)

    def para(self, text, indent=True):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.125)
        # **bold** segments
        parts = text.split("**")
        for i, part in enumerate(parts):
            if part:
                self._math_runs(p, part, bold=(i % 2 == 1))
        return p

    def bullets(self, items):
        for item in items:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.12)
            parts = ("• " + item).split("**")
            for i, part in enumerate(parts):
                if part:
                    self._run(p, part, bold=(i % 2 == 1))

    def equation(self, text):
        self.equation_no += 1
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        stops = p.paragraph_format.tab_stops
        stops.add_tab_stop(Emu(COLUMN_WIDTH // 2), WD_TAB_ALIGNMENT.CENTER)
        stops.add_tab_stop(COLUMN_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
        self._run(p, "\t")
        self._math_runs(p, text)
        self._run(p, f"\t({self.equation_no})")

    def figure(self, filename, caption, width=COLUMN_WIDTH):
        self.figure_no += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(os.path.join(FIG, filename), width=width)
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        c.paragraph_format.space_before = Pt(2)
        c.paragraph_format.space_after = Pt(8)
        self._run(c, f"Fig. {self.figure_no}. ", size=8)
        self._run(c, caption, size=8)
        return self.figure_no

    def _set_columns(self, section, number):
        cols = section._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            section._sectPr.append(cols)
        cols.set(qn("w:num"), str(number))
        cols.set(qn("w:space"), "360")

    def wide_figure(self, filename, caption):
        """Figure spanning both columns (one-column section between two-column ones)."""
        self._set_columns(self.doc.add_section(WD_SECTION.CONTINUOUS), 1)
        number = self.figure(filename, caption, width=Inches(6.6))
        self._set_columns(self.doc.add_section(WD_SECTION.CONTINUOUS), 2)
        return number

    def table(self, caption, header, rows, widths=None, bold_rows=()):
        self.table_no += 1
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_before = Pt(6)
        c.paragraph_format.keep_with_next = True
        self._run(c, f"TABLE {ROMAN[self.table_no - 1]}", size=8)
        c2 = self.doc.add_paragraph()
        c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c2.paragraph_format.space_after = Pt(3)
        c2.paragraph_format.keep_with_next = True
        self._run(c2, caption, size=8, small_caps=True)

        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        table._tbl.tblPr.append(layout)
        margins = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), "45")
            m.set(qn("w:type"), "dxa")
            margins.append(m)
        table._tbl.tblPr.append(margins)
        if widths:
            for k, width in enumerate(widths):
                table.columns[k].width = width
        for r, values in enumerate([header] + rows):
            for k, value in enumerate(values):
                cell = table.rows[r].cells[k]
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if k else WD_ALIGN_PARAGRAPH.LEFT
                self._run(cell.paragraphs[0], str(value), size=7.5, bold=(r == 0 or (r - 1) in bold_rows))
                if widths:
                    cell.width = widths[k]
                # Keep every row with the next so the table never splits across pages
                if r < len(rows):
                    cell.paragraphs[0].paragraph_format.keep_with_next = True
            row_props = table.rows[r]._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            row_props.append(cant_split)
        after = self.doc.add_paragraph()
        after.paragraph_format.space_after = Pt(6)
        return self.table_no

    def references(self, refs):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        self._run(p, "References", small_caps=True)
        for i, ref in enumerate(refs, 1):
            r = self.doc.add_paragraph()
            r.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r.paragraph_format.left_indent = Inches(0.3)
            r.paragraph_format.first_line_indent = Inches(-0.3)
            r.paragraph_format.space_after = Pt(1)
            self._run(r, f"[{i}] {ref}", size=8)

    def _hyperlink(self, paragraph, url, size=9):
        """Clickable link run (python-docx has no direct API for hyperlinks)."""
        r_id = paragraph.part.relate_to(
            url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        font_size = OxmlElement("w:sz")
        font_size.set(qn("w:val"), str(int(size * 2)))
        props.extend([color, underline, font_size])
        text = OxmlElement("w:t")
        text.text = url
        run.extend([props, text])
        link.append(run)
        paragraph._p.append(link)

    def appendix(self, title, intro, items):
        """Unnumbered appendix with labelled links, placed after the references."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        self._run(p, title, small_caps=True)
        self.para(intro)
        for label, url in items:
            item = self.doc.add_paragraph()
            item.paragraph_format.space_before = Pt(3)
            item.paragraph_format.left_indent = Inches(0.1)
            self._run(item, f"{label}: ", size=9, bold=True)
            self._hyperlink(item, url)

    def save(self):
        self.doc.save(OUT)
        return OUT


def build():
    p = Paper()
    p.title_block()
    p.start_two_columns()

    p.abstract(
        "Microservice platforms emit large volumes of logs, and operators need to notice failures "
        "quickly and understand where they started. This paper presents a log monitoring system built "
        "on MongoDB that summarizes each service's logs into overlapping five-minute windows of 19 "
        "features and scores every window with two complementary detectors: a severity-weighted rule "
        "engine and an unsupervised Isolation Forest. The two confidences are fused with fixed weights "
        "(0.4 and 0.6) into a single decision. When an incident occurs, an investigator correlates "
        "anomalies across services, reconstructs a minute-level timeline from the raw logs, ranks "
        "root-cause hypotheses with a service dependency map, and asks a locally hosted large language "
        "model to write a report that must cite numbered evidence; citations are verified and "
        "state-changing advice is flagged for human approval. On five labelled runs (5,272 windows), the "
        "hybrid detector reaches an F1-score of 0.870 and accuracy of 0.983 without any labels, ahead of "
        "Local Outlier Factor (0.371) and One-Class SVM (0.608). A supervised Random Forest scores higher "
        "(F1 0.931) when it is trained on 4,218 labelled windows of every incident type, but it detects "
        "none of the error bursts or traffic drops once that incident type is absent from its training "
        "data, whereas our Isolation Forest still recalls 84-85% of them. The investigator identified the "
        "correct origin of a three-service cascading failure in five of five runs.",
        "anomaly detection, log analysis, isolation forest, random forest, MongoDB, microservices, "
        "root cause analysis, large language models")

    # I. Introduction
    p.heading("Introduction")
    p.para("Modern applications are split into many small services that call one another over the "
           "network. A fault in one service rarely stays inside it: a slow database behind a user "
           "service makes the authentication service time out, which in turn makes payments fail. Each "
           "service writes its own log stream, so an operator who looks at a failing payment service "
           "sees only symptoms. Two questions therefore matter during an incident: whether something is "
           "wrong right now, and which component started the problem.")
    p.para("Fixed threshold alerts answer the first question cheaply and are easy to interpret, but they "
           "only catch conditions that someone anticipated. Machine-learning detectors can flag unusual "
           "behaviour that no rule describes, yet supervised models need labelled examples of failures, "
           "which are scarce and quickly outdated, and unsupervised models tend to raise false alarms. "
           "Neither kind of detector explains where a multi-service failure began.")
    p.para("This work combines both detectors and adds an evidence-driven investigation step. Its "
           "contributions are:")
    p.bullets([
        "A MongoDB data design for log monitoring that keeps raw logs, window features, detection "
        "results, rules, versioned models (in GridFS) and incident reports in one store, using compound "
        "indexes, aggregation pipelines and time-to-live expiry.",
        "A hybrid detector that fuses a severity-weighted rule confidence with a normalized Isolation "
        "Forest score, needs no labelled data, and explains every alert with the rules, values and "
        "thresholds involved.",
        "An incident investigator that correlates anomalies across services, finds each service's onset "
        "from raw logs, scores root-cause hypotheses with a dependency map, and grounds a local LLM's "
        "report in verified evidence without allowing any automatic action.",
        "An evaluation against Random Forest, Local Outlier Factor and One-Class SVM, including a test on "
        "incident types that were never seen during training.",
    ])

    # II. Related work
    p.heading("Related Work")
    p.para("Early log anomaly detectors treated each session as a vector of event counts. Xu et al. "
           "applied principal component analysis to console logs [5], and Lou et al. mined linear "
           "invariants between event counts [6]. He et al. compared these and several supervised methods "
           "on public datasets and released the loglizer toolkit [7]; a later survey summarizes the "
           "field [14]. Deep models then learned event sequences: DeepLog predicts the next log key with "
           "an LSTM [8], LogAnomaly adds semantic templates [9], and LogBERT trains a transformer with "
           "self-supervised objectives [10]. Le and Zhang showed that the reported accuracy of deep "
           "methods drops under realistic data splits [11]. These approaches focus on the order of log "
           "messages; they do not use response time or traffic volume, and they output a binary label "
           "without explaining it.")
    p.para("Isolation Forest isolates outliers with random partitions and needs no labels [1]. Random "
           "Forest is a strong supervised ensemble [2], while Local Outlier Factor [3] and One-Class SVM "
           "[4] are common unsupervised baselines. Recent studies use large language models to suggest "
           "root causes and mitigations for cloud incidents [12], [13]. Language models can, however, "
           "state facts that are not supported by their input [17], which is risky during an outage. Our "
           "investigator therefore computes the root cause from data first and only lets the model "
           "explain it, with every statement checked against the evidence.")

    # III. Architecture
    p.heading("System Architecture")
    p.figure("fig_architecture.png", "Data flow of the system, from log ingestion to the dashboard.")
    p.para("Fig. 1 shows the pipeline. Services send logs to a FastAPI backend through single or batch "
           "ingestion endpoints. Each log carries a timestamp, service name, level, message, endpoint, "
           "HTTP method, status code, response time, payload sizes and a free-form metadata field. Logs "
           "are stored in MongoDB, summarized into windows, scored by the two detectors, and fused into a "
           "decision that is stored together with its explanation. A React dashboard reads the results "
           "through REST endpoints, and the investigator is invoked from the dashboard for a selected "
           "anomaly.")
    p.subheading("MongoDB Data Design")
    p.para("The database holds seven collections. **log_entries** stores raw logs; its flexible "
           "document schema lets each service attach different metadata without migrations. "
           "**window_features** stores one document per service and window with the 19-value feature "
           "vector. **anomaly_results** stores every decision with the rule, model and fused "
           "confidences and the details of each violated rule. **detection_rules** keeps the rule "
           "definitions as data, so thresholds can change without code changes. **model_metadata** "
           "records model versions and parameters, while the serialized model itself is stored in "
           "**GridFS** because it exceeds the size intended for ordinary documents. "
           "**incident_reports** keeps investigation results.")
    p.para("Compound indexes on (service_name, timestamp) and (window_start, service_name) serve the "
           "dashboard's most frequent query, the latest records of one service. Time-to-live indexes "
           "remove logs after 30 days, windows after 90 days and results after 180 days, so storage "
           "stays bounded without a cleanup job. Aggregation pipelines compute averages, the 95th "
           "percentile latency with $percentile, error counts and per-minute timelines inside the "
           "database rather than in application code.")

    # IV. Methodology
    p.heading("Methodology")
    p.subheading("Windowing and Features")
    p.para("Logs of each service are grouped into windows of five minutes that start every minute, so "
           "consecutive windows overlap and a problem is seen within about a minute. A window is scored "
           "only if it contains at least five requests. Table I lists the features computed per window.")
    p.table("Window features", ["Group", "Features"], [
        ["Volume", "request count, error count approximation"],
        ["Errors", "error rate (ERROR logs / requests)"],
        ["Latency", "mean, 95th and 99th percentile, maximum, time per request"],
        ["Payload", "mean request size, mean response size"],
        ["HTTP methods", "GET, POST, PUT, DELETE, PATCH counts"],
        ["Status codes", "2xx, 3xx, 4xx, 5xx counts"],
    ], widths=[Inches(0.9), Inches(2.5)])

    p.subheading("Rule Engine and Rule Confidence")
    p.para("Seven rules encode operational limits (Table II). Threshold rules compare a feature with a "
           "fixed limit. Rate-change rules compare it with the mean of the previous three windows; for "
           "error rate this mean is floored at 5% so that a change from zero to one error is not "
           "reported as a large relative spike.")
    p.table("Detection rules", ["Rule", "Condition", "Severity"], [
        ["high_error_rate", "error rate > 10%", "high"],
        ["very_high_error_rate", "error rate > 25%", "critical"],
        ["high_latency_p95", "P95 > 1000 ms", "medium"],
        ["high_latency_p99", "P99 > 2000 ms", "high"],
        ["low_traffic", "requests < 10", "medium"],
        ["sudden_traffic_drop", "requests < 0.5 x recent", "high"],
        ["sudden_error_spike", "error rate > 3 x recent", "high"],
    ], widths=[Inches(1.35), Inches(1.45), Inches(0.6)])
    p.para("Each violated rule i receives a score s_i in [0, 1]. For threshold rules the score grows "
           "with the relative distance beyond the limit, s_i = min(1, |v - t| / t), where v is the "
           "observed value and t the threshold; rate-change rules receive s_i = 0.8. Each rule also has a "
           "severity weight w_i (critical 1.0, high 0.8, medium 0.6). The rule confidence is the "
           "weighted mean")
    p.equation("R = Σ w_i s_i / Σ w_i")
    p.para("and R = 0 when no rule fires. For example, a window with 81.8% errors and P95 of 4940 ms "
           "violates six rules; four threshold rules score 1.0 and two rate-change rules score 0.8, which "
           "gives R = 4.48 / 4.8 = 0.93.")

    p.subheading("Isolation Forest and Model Confidence")
    p.para("Isolation Forest builds random binary trees in which every node splits on a randomly "
           "chosen feature at a random value [1]. Points that differ from the bulk of the data are "
           "separated after few splits, so a short average path length indicates an anomaly. We train "
           "100 trees, each on a random 80% of the windows, after standardizing every feature to zero "
           "mean and unit variance so that millisecond-scale latency does not dominate rates between 0 "
           "and 1. The contamination parameter is set to 0.1, which places the decision boundary at the "
           "most isolated 10% of training windows. Training uses no labels and takes well under a "
           "second; the fitted model and scaler are versioned and stored in GridFS.")
    p.para("The library's decision score d is zero at the boundary and negative for anomalies. We map it "
           "to a confidence in [0, 1] with a steep logistic function")
    p.equation("M = 1 / (1 + e^(20 d))")
    p.para("so that the boundary corresponds to M = 0.5 and typical scores between -0.2 and 0.2 spread "
           "across the whole range. With a slope of one, almost all normal windows had confidences "
           "between 0.48 and 0.50 and were mislabelled as borderline. A window with 0.4 <= M < 0.5 is "
           "borderline; three consecutive borderline windows raise an alert, which catches slow "
           "degradations that never look extreme in one window.")

    p.subheading("Hybrid Fusion")
    p.para("The two confidences are combined as")
    p.equation("F = 0.4 R + 0.6 M,  anomaly if F ≥ 0.5")
    p.para("The model receives the larger weight because it generalizes to patterns no rule describes. "
           "With these weights, strong rule evidence (R = 1) needs only M ≥ 0.17 to alert, while "
           "the model alone needs M ≥ 0.84, which suppresses many of its false alarms. The fused "
           "decision is labelled hybrid_both when both detectors flag the window, rule_only or "
           "model_only otherwise, and borderline_persistence when the persistence rule fired. Until a "
           "model exists the system falls back to rules alone. Repeated alerts for the same service and "
           "the same violated rules within five minutes of window time are stored but not raised again.")

    p.subheading("Real-Time and Batch Detection")
    p.para("When logs arrive through the API, the backend builds the newest window of that service from "
           "the last ten minutes, scores it, and stores both the window and the decision; a decision "
           "takes 11 to 25 ms. A batch pipeline processes all stored logs in time order, trains the "
           "model if none is stored, and is exposed through a command-line tool and an API endpoint.")

    p.subheading("Evidence-Grounded Incident Investigation")
    p.para("The investigator starts from one anomaly and collects the anomalies of all services within "
           "15 minutes before and after it. For each affected service it computes a baseline from the 60 "
           "windows preceding the incident and scans the raw logs minute by minute. A minute is "
           "abnormal if it has at least three errors and an error rate of at least max(30%, 5x "
           "baseline), or a mean latency of at least three times the baseline. The onset is the first "
           "minute of a run of at least two abnormal minutes, which prevents a single noisy minute from "
           "being taken as the start.")
    p.para("Findings are recorded as numbered evidence (E1, E2, ...): anomaly records, metrics compared "
           "with the baseline, the onset minute, and the most frequent error messages, each pointing to "
           "its MongoDB source. Every affected service c is then scored as a root-cause candidate:")
    p.equation("S_c = 0.3 C_c + 0.25 f_c + 0.25 D_c + 0.2 m_c")
    p.para("where C_c is the highest anomaly confidence of c, f_c is 1 if c degraded first, D_c is the "
           "fraction of other affected services that depend on c and degraded after it, and m_c is 1 if "
           "error messages of other services name c. A single affected service receives 0.2 more, and "
           "scores are capped at 0.99. Severity, read-only troubleshooting steps and a timeline are "
           "derived from the same evidence.")
    p.para("The computed conclusion, alternatives, timeline and evidence are sent to a local model "
           "(llama3.2 through Ollama, temperature 0.1) with instructions to cite evidence identifiers in "
           "every sentence and never to recommend state-changing actions. The response is checked: "
           "citations to unknown identifiers and sentences without citations are listed for the reader, "
           "and words such as restart, delete or roll back are flagged as requiring human approval. The "
           "investigator only reads data and the model has no tools, so nothing is executed "
           "automatically. If the model is unavailable, the data-driven part of the report is still "
           "produced. Running the model locally keeps logs inside the organization and avoids API cost.")

    # V. Implementation
    p.heading("Implementation")
    p.para("The backend uses Python, FastAPI, PyMongo and scikit-learn [16]; the dashboard uses React, "
           "Material UI and Recharts. The dashboard shows system totals, per-service latency and error "
           "metrics, trend charts, and tables of anomalies and logs (Fig. 2). The details view of an "
           "anomaly lists the fused, rule and model confidences and every violated rule with its value, "
           "threshold and severity (Fig. 3). A button starts the investigation (Fig. 6), and another "
           "simulates a cascading failure in which the user service's database slows down, the "
           "authentication service then times out calling it, and the payment service fails one minute "
           "later.")
    p.figure("dashboard.png", "Monitoring dashboard with system totals, per-service metrics and the "
             "breakdown of detection methods.")
    p.figure("anomaly_details.png", "Explanation of one alert: fused confidence 0.76 = 0.4 x 0.67 + "
             "0.6 x 0.82, with the violated rule, observed value and threshold.")

    # VI. Experimental setup
    p.heading("Experimental Setup")
    p.para("Labelled anomalies are not available for real microservice traffic, and public log "
           "benchmarks such as HDFS record event sequences without response times. We therefore "
           "generated labelled traffic with a seeded generator. Each run covers six hours for three "
           "services at about eight requests per minute (about 8,000 logs, 1,054 windows) with "
           "log-normal latencies around 135 ms and 2% background errors. Four incidents are injected: a "
           "15-minute latency spike in the payment service, a 30-minute gradual slowdown and a "
           "10-minute error burst in the authentication service, and a 15-minute traffic drop in the "
           "user service. A window is labelled anomalous if it contains logs of an injected incident. "
           "Five runs with different seeds give 5,272 windows, of which 375 are anomalous.")
    p.para("The unsupervised methods (our rule engine, Isolation Forest and hybrid, Local Outlier Factor "
           "with 20 neighbours and One-Class SVM with nu = 0.1) are fitted on the unlabelled windows "
           "they monitor. Random Forest (100 trees) is supervised: for each run it is trained on the "
           "labelled windows of the other four runs (4,218 windows) and tested on the held-out run. "
           "Experiment A gives Random Forest examples of every incident type. Experiment B removes one "
           "incident type from its training data and measures recall on that type, which models a "
           "failure mode that has not occurred before. Precision, recall, F1-score, accuracy and false "
           "alarms are summed over the five runs. The scripts evaluation/model_comparison.py and "
           "evaluation/investigation_eval.py reproduce all numbers.")

    # VII. Results
    p.heading("Results and Discussion")
    p.subheading("Comparison with Random Forest and Unsupervised Baselines")
    p.table("Detection results over five runs (Experiment A); Alarms = false alarms",
            ["Model", "Labels", "Prec.", "Rec.", "F1", "Acc.", "Alarms"], [
                ["Rules (ours)", "no", "0.901", "0.752", "0.820", "0.976", "31"],
                ["IF (ours)", "no", "0.638", "0.901", "0.747", "0.957", "192"],
                ["Hybrid (ours)", "no", "0.950", "0.803", "0.870", "0.983", "16"],
                ["Random Forest", "yes", "0.963", "0.901", "0.931", "0.991", "13"],
                ["LOF", "no", "0.317", "0.448", "0.371", "0.892", "362"],
                ["One-Class SVM", "no", "0.521", "0.731", "0.608", "0.933", "252"],
            ], widths=[Inches(0.95), Inches(0.4), Inches(0.42), Inches(0.42), Inches(0.42), Inches(0.42),
                       Inches(0.42)], bold_rows=(2,))
    p.figure("fig_known.png", "Precision, recall and F1-score of each model in Experiment A. "
             "*Random Forest is the only model trained with labels.")
    p.para("Table III shows that fusion improves on both of its parts. The rule engine is precise but "
           "misses a quarter of the anomalous windows, and the Isolation Forest finds 90% of them at the "
           "cost of 192 false alarms. The hybrid keeps most of the rules' precision (0.950) while "
           "raising recall to 0.803, giving an F1-score of 0.870 and only 16 false alarms in 4,897 "
           "normal windows. Among methods that use no labels it is clearly the strongest: Local Outlier "
           "Factor and One-Class SVM reach F1-scores of 0.371 and 0.608 and raise 15 to 22 times more "
           "false alarms.")
    p.para("Random Forest achieves the best scores (F1 0.931, accuracy 0.991). This is expected: it is "
           "trained on 4,218 labelled windows that contain the same four incident types that appear in "
           "the test run. Accuracy is high for every model because 93% of windows are normal; F1-score "
           "and false alarms are more informative.")

    p.subheading("Detecting Incident Types Not Seen in Training")
    p.table("Recall on an incident type absent from Random Forest training (Experiment B)",
            ["Unseen type", "Random Forest", "IF (ours)", "Hybrid (ours)"], [
                ["Latency spike", "1.000", "1.000", "1.000"],
                ["Gradual slowdown", "0.842", "0.877", "0.702"],
                ["Error burst", "0.000", "0.849", "0.822"],
                ["Traffic drop", "0.000", "0.844", "0.688"],
                ["Mean", "0.461", "0.893", "0.803"],
            ], widths=[Inches(1.1), Inches(0.8), Inches(0.7), Inches(0.8)], bold_rows=(4,))
    p.figure("fig_unseen.png", "Recall on an incident type that Random Forest never saw during "
             "training. Our detectors use no labels, so their results do not depend on which types "
             "occurred before.")
    p.para("Table IV changes the picture. When error bursts are missing from its training data, Random "
           "Forest detects none of them, and the same happens for traffic drops: it learned what the "
           "known incidents look like rather than what normal traffic looks like. Latency spikes and "
           "slowdowns are still caught because they resemble each other. Our Isolation Forest models "
           "normal behaviour only, so it recalls 84% to 100% of every type, 0.893 on average against "
           "0.461 for Random Forest. The hybrid recalls 0.803 on average while keeping far fewer false "
           "alarms than the Isolation Forest alone.")

    p.subheading("Why the Hybrid Approach Is Preferable in Practice")
    p.para("The comparison supports the design for operational use. First, it needs no labels: Random "
           "Forest required 4,218 labelled windows per run, which a new deployment does not have. "
           "Second, it keeps detecting failure modes that have never been observed, where the "
           "supervised model's recall dropped to zero. Third, every alert is explained by the rules, "
           "values and thresholds that caused it, while a Random Forest vote does not tell an operator "
           "what is wrong. Fourth, it adapts by retraining on recent unlabelled data in under a second. "
           "When large labelled histories exist, a supervised model can be added as a further input to "
           "the fusion step.")

    p.subheading("Incident Detection and Investigation")
    p.para("In every run all four injected incidents produced alerts; the gradual slowdown was partly "
           "caught by model_only alerts that no rule raised. For the cascading failure, the investigator "
           "was run five times with different background traffic. In all five runs it ranked the user "
           "service as the root cause with 97-99% confidence and reconstructed the correct order of "
           "failure (user, then authentication, then payment service). Fig. 6 shows one report: the "
           "user service degraded first, both dependent services failed afterwards, and their error "
           "messages name it. An earlier version took a single noisy minute as the onset and blamed the "
           "authentication service; requiring two consecutive abnormal minutes removed this error.")
    p.figure("investigation.png", "Incident investigation of a simulated cascading failure: ranked "
             "root-cause hypotheses with evidence identifiers and the reconstructed timeline.")
    p.para("With the 3-billion-parameter llama3.2 model, reports followed the requested structure but "
           "sometimes omitted citations or suggested generic actions such as restarting services. The "
           "validation step listed the uncited sentences and flagged the restart suggestion for human "
           "approval, which illustrates why the conclusion is computed from data and the language "
           "model is used only to explain it.")

    # VIII. Limitations
    p.heading("Limitations and Future Work")
    p.para("The evaluation uses generated traffic with injected incidents, so absolute scores on "
           "production systems may differ; the relative behaviour of supervised and unsupervised models "
           "on unseen incident types is the main finding. Window labels mark any window that overlaps "
           "an incident, which slightly lowers the measured recall at incident boundaries. The "
           "dependency map is configured by hand, the dashboard polls every 30 seconds, and the API has "
           "no authentication. Future work includes learning dependencies from traces, operator "
           "feedback on alerts, webhook notifications, and evaluation on public microservice "
           "benchmarks that include latency.")

    # IX. Conclusion
    p.heading("Conclusion")
    p.para("We presented a MongoDB-based system that detects anomalies in microservice logs by fusing a "
           "severity-weighted rule engine with an unsupervised Isolation Forest and that investigates "
           "incidents with an evidence-grounded local language model. Without labels, the hybrid "
           "detector reached an F1-score of 0.870 with 16 false alarms in about 4,900 normal windows, "
           "well ahead of other unsupervised methods. A supervised Random Forest scored higher only when "
           "it had been trained on labelled examples of every incident type, and it missed all error "
           "bursts and traffic drops that were absent from its training data, while our Isolation Forest "
           "still detected 84-85% of them. The investigator located the origin of a cascading failure in "
           "every trial and never executes actions, which makes the approach suitable for environments "
           "where failure labels are rare and explanations are required.")

    p.references([
        "F. T. Liu, K. M. Ting, and Z.-H. Zhou, “Isolation forest,” in Proc. IEEE Int. Conf. "
        "Data Mining (ICDM), 2008, pp. 413–422.",
        "L. Breiman, “Random forests,” Machine Learning, vol. 45, no. 1, pp. 5–32, 2001.",
        "M. M. Breunig, H.-P. Kriegel, R. T. Ng, and J. Sander, “LOF: Identifying density-based local "
        "outliers,” in Proc. ACM SIGMOD Int. Conf. Management of Data, 2000, pp. 93–104.",
        "B. Schölkopf, J. C. Platt, J. Shawe-Taylor, A. J. Smola, and R. C. Williamson, “Estimating "
        "the support of a high-dimensional distribution,” Neural Computation, vol. 13, no. 7, "
        "pp. 1443–1471, 2001.",
        "W. Xu, L. Huang, A. Fox, D. Patterson, and M. I. Jordan, “Detecting large-scale system "
        "problems by mining console logs,” in Proc. ACM Symp. Operating Systems Principles (SOSP), "
        "2009, pp. 117–132.",
        "J.-G. Lou, Q. Fu, S. Yang, Y. Xu, and J. Li, “Mining invariants from console logs for system "
        "problem detection,” in Proc. USENIX Annual Technical Conf., 2010.",
        "S. He, J. Zhu, P. He, and M. R. Lyu, “Experience report: System log analysis for anomaly "
        "detection,” in Proc. IEEE Int. Symp. Software Reliability Engineering (ISSRE), 2016, "
        "pp. 207–218.",
        "M. Du, F. Li, G. Zheng, and V. Srikumar, “DeepLog: Anomaly detection and diagnosis from system "
        "logs through deep learning,” in Proc. ACM Conf. Computer and Communications Security (CCS), "
        "2017, pp. 1285–1298.",
        "W. Meng et al., “LogAnomaly: Unsupervised detection of sequential and quantitative anomalies "
        "in unstructured logs,” in Proc. Int. Joint Conf. Artificial Intelligence (IJCAI), 2019, "
        "pp. 4739–4745.",
        "H. Guo, S. Yuan, and X. Wu, “LogBERT: Log anomaly detection via BERT,” in Proc. Int. "
        "Joint Conf. Neural Networks (IJCNN), 2021.",
        "V.-H. Le and H. Zhang, “Log-based anomaly detection with deep learning: How far are we?” "
        "in Proc. IEEE/ACM Int. Conf. Software Engineering (ICSE), 2022, pp. 1356–1367.",
        "T. Ahmed, S. Ghosh, C. Bansal, T. Zimmermann, X. Zhang, and S. Rajmohan, “Recommending "
        "root-cause and mitigation steps for cloud incidents using large language models,” in Proc. "
        "IEEE/ACM Int. Conf. Software Engineering (ICSE), 2023.",
        "Y. Chen et al., “Automatic root cause analysis via large language models for cloud "
        "incidents,” in Proc. European Conf. Computer Systems (EuroSys), 2024.",
        "S. He, P. He, Z. Chen, T. Yang, Y. Su, and M. R. Lyu, “A survey on automated log analysis "
        "for reliability engineering,” ACM Computing Surveys, vol. 54, no. 6, 2021.",
        "K. Banker, P. Bakkum, S. Verch, D. Garrett, and T. Hawkins, MongoDB in Action, 2nd ed. Shelter "
        "Island, NY, USA: Manning, 2016.",
        "F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” Journal of Machine "
        "Learning Research, vol. 12, pp. 2825–2830, 2011.",
        "Z. Ji et al., “Survey of hallucination in natural language generation,” ACM Computing "
        "Surveys, vol. 55, no. 12, 2023.",
    ])
    p.appendix(
        "Appendix: Drive Link",
        "The project folder contains the complete source code, the demonstration video of the "
        "working system from log ingestion to AI incident investigation, and the user manual that "
        "explains how to install and run it. The folder is shared so that anyone with the link can "
        "view it.",
        [
            ("Drive folder (code, demo video, user manual)",
             "https://drive.google.com/drive/folders/1YvUnh0IvpU9PN5cbIb2BNfOtxddi_WOl"),
            ("Source code repository", "https://github.com/sivakumar1306/Nosql-finalyear-project"),
        ])
    return p.save()


if __name__ == "__main__":
    print(build())
